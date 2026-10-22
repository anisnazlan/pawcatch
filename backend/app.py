"""WildWatch Elmina backend: resident reports, AI classification, routing and dashboard API.

Run:  uvicorn backend.app:app --reload   (from the project root)

Reports are saved in data/wildwatch.db (SQLite) and photos in data/uploads/.
Settings (environment variables):
  ADMIN_KEY      dashboard login key (default: elmina2026)
  SHEETS_URL     optional: Google Apps Script web app link, for a permanent copy of reports
  SHEETS_SECRET  optional: the secret written in google_sheets/Code.gs
"""
import os
import re
import uuid
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Optional

from fastapi import FastAPI, File, Form, Header, HTTPException, UploadFile
from fastapi.responses import FileResponse, RedirectResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

from . import rules, sheets
from .classifier import classify, model_ready
from .db import init_db, query

ROOT = Path(__file__).resolve().parent.parent
UPLOADS = ROOT / "data" / "uploads"
UPLOADS.mkdir(parents=True, exist_ok=True)
ADMIN_KEY = os.getenv("ADMIN_KEY", "elmina2026")
STATUSES = ["received", "assigned", "in_progress", "resolved"]
IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".webp", ".heic", ".heif"}
MAX_PHOTO_BYTES = 10 * 1024 * 1024

app = FastAPI(title="WildWatch Elmina")
init_db()


def restore_from_sheet():
    """After a restart, put live reports from the Google Sheet back, with their original IDs."""
    rows = sheets.load_reports()
    added = 0
    for r in rows:
        n = query(
            """INSERT OR IGNORE INTO reports (id, created_at, lat, lng, animal, incident, urgency,
               agency, status, confidence, ai_mode, photo, note, resolved_at, simulated)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,NULL,?,?,0)""",
            (r["id"], r["created_at"], r["lat"], r["lng"], r["animal"], r["incident"], r["urgency"],
             r["agency"], r["status"], r["confidence"], r["ai_mode"], r["note"], r["resolved_at"]),
            fetch="none")
        added += n or 0
    if sheets.ENABLED:
        print(f"[sheets] restored {added} live report(s) from the Google Sheet", flush=True)


restore_from_sheet()


@app.on_event("startup")
def show_links():
    # Clickable in the VS Code terminal: Ctrl + click to open.
    if os.getenv("RENDER"):
        return
    port = os.getenv("PORT", "8000")
    print(f"\n  WildWatch is running. Ctrl + click a link to open it:\n"
          f"    Resident app: http://localhost:{port}/app\n"
          f"    Dashboard:    http://localhost:{port}/dashboard\n", flush=True)


def label(d: dict) -> dict:
    d["animal_label"] = rules.ANIMAL_LABELS.get(d["animal"], d["animal"])
    d["incident_label"] = rules.INCIDENT_LABELS.get(d["incident"], d["incident"])
    d["agency_label"] = rules.AGENCIES.get(d["agency"], d["agency"])
    return d


def valid_photo_ref(ref: Optional[str]) -> Optional[str]:
    """Only accept photo references this server created."""
    if not ref:
        return None
    return ref if re.fullmatch(r"[0-9a-f]{32}\.[a-z]{3,4}", ref) else None


# ---------- Resident side ----------

@app.post("/api/classify")
def classify_photo(photo: UploadFile = File(...)):
    """Step 1: resident uploads a photo, AI suggests the animal, photo is stored."""
    ext = Path(photo.filename or "img.jpg").suffix.lower()
    if ext not in IMAGE_EXTS:
        ext = ".jpg"
    data = photo.file.read(MAX_PHOTO_BYTES + 1)
    if len(data) > MAX_PHOTO_BYTES:
        raise HTTPException(413, "Photo is larger than 10 MB")
    path = UPLOADS / f"{uuid.uuid4().hex}{ext}"
    path.write_bytes(data)
    try:
        result = classify(str(path))
    except Exception as e:  # an unreadable image should not block the report
        print(f"[classifier] failed: {e}")
        result = {"animal": None, "confidence": 0.0, "confident": False, "mode": "error"}
    result["photo"] = path.name
    return result


@app.post("/api/reports")
def create_report(
    lat: float = Form(...),
    lng: float = Form(...),
    animal: str = Form(...),
    incident: str = Form(...),
    photo: Optional[str] = Form(None),
    confidence: Optional[float] = Form(None),
    ai_mode: Optional[str] = Form(None),
    note: Optional[str] = Form(None),
):
    """Step 2: resident confirms animal and incident, system routes the report."""
    if animal not in rules.ANIMALS or incident not in rules.INCIDENTS:
        raise HTTPException(400, "Unknown animal or incident type")
    if not (-90 <= lat <= 90 and -180 <= lng <= 180):
        raise HTTPException(400, "Invalid location")
    urgency = rules.urgency_for(animal, incident)
    agency = rules.agency_for(animal, incident)
    now = datetime.now().isoformat(timespec="seconds")
    row = query(
        """INSERT INTO reports (created_at, lat, lng, animal, incident, urgency, agency,
           status, confidence, ai_mode, photo, note) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)
           RETURNING id""",
        (now, lat, lng, animal, incident, urgency, agency, "assigned",
         confidence, (ai_mode or "")[:20] or None, valid_photo_ref(photo), (note or "")[:500] or None),
        fetch="one",
    )
    sheets.add_report(query("SELECT * FROM reports WHERE id=?", (row["id"],), fetch="one"))
    return {
        "id": row["id"],
        "animal_label": rules.ANIMAL_LABELS[animal],
        "urgency": urgency,
        "agency": agency,
        "agency_label": rules.AGENCIES[agency],
        "tips": rules.TIPS[animal],
        "status": "assigned",
    }


@app.get("/api/reports/{rid}")
def get_report(rid: int):
    r = query("SELECT * FROM reports WHERE id=?", (rid,), fetch="one")
    if not r:
        raise HTTPException(404, "Report not found")
    d = label(r)
    d["tips"] = rules.TIPS.get(d["animal"], [])
    return d


# ---------- Management side ----------

def require_admin(key: Optional[str]):
    if key != ADMIN_KEY:
        raise HTTPException(401, "Invalid admin key")


class LoginBody(BaseModel):
    key: str


@app.post("/api/login")
def login(body: LoginBody):
    require_admin(body.key)
    return {"ok": True}


class StatusBody(BaseModel):
    status: str


@app.patch("/api/reports/{rid}")
def update_status(rid: int, body: StatusBody, x_admin_key: Optional[str] = Header(None)):
    require_admin(x_admin_key)
    if body.status not in STATUSES:
        raise HTTPException(400, "Unknown status")
    resolved = datetime.now().isoformat(timespec="seconds") if body.status == "resolved" else None
    n = query("UPDATE reports SET status=?, resolved_at=COALESCE(?, resolved_at) WHERE id=?",
              (body.status, resolved, rid), fetch="none")
    if not n:
        raise HTTPException(404, "Report not found")
    r = query("SELECT simulated, resolved_at FROM reports WHERE id=?", (rid,), fetch="one")
    if not r["simulated"]:
        sheets.update_status(rid, body.status, r["resolved_at"])
    return {"ok": True}


@app.get("/api/reports")
def list_reports(x_admin_key: Optional[str] = Header(None), limit: int = 2000):
    require_admin(x_admin_key)
    rows = query("SELECT * FROM reports ORDER BY created_at DESC LIMIT ?", (min(limit, 5000),))
    return [label(r) for r in rows]


def hotspots(rows, cell=0.0015, top=5):
    """Grid-based hotspot detection: roughly 165 m cells, ranked by report count."""
    cells = defaultdict(list)
    for r in rows:
        cells[(round(r["lat"] / cell), round(r["lng"] / cell))].append(r)
    ranked = sorted(cells.values(), key=len, reverse=True)[:top]
    out = []
    for group in ranked:
        animal = Counter(g["animal"] for g in group).most_common(1)[0][0]
        incident = Counter(g["incident"] for g in group if g["animal"] == animal).most_common(1)[0][0]
        hours = [datetime.fromisoformat(g["created_at"]).hour for g in group if g["animal"] == animal]
        period = Counter(rules.period_of(h) for h in hours).most_common(1)[0][0]
        out.append({
            "lat": sum(g["lat"] for g in group) / len(group),
            "lng": sum(g["lng"] for g in group) / len(group),
            "count": len(group),
            "animal": animal,
            "animal_label": rules.ANIMAL_LABELS[animal],
            "incident_label": rules.INCIDENT_LABELS[incident],
            "period": period,
            "suggestion": rules.suggestion_for(animal, incident, period),
        })
    return out


@app.get("/api/stats")
def stats(x_admin_key: Optional[str] = Header(None)):
    require_admin(x_admin_key)
    rows = query("SELECT * FROM reports")
    now = datetime.now()
    month = [r for r in rows if r["created_at"][:7] == now.strftime("%Y-%m")]
    res_hours = []
    for r in rows:
        if r["resolved_at"]:
            delta = datetime.fromisoformat(r["resolved_at"]) - datetime.fromisoformat(r["created_at"])
            res_hours.append(delta.total_seconds() / 3600)
    animals = Counter(r["animal"] for r in rows)
    return {
        "total": len(rows),
        "this_month": len(month),
        "urgent_open": sum(1 for r in rows if r["urgency"] == "high" and r["status"] != "resolved"),
        "avg_resolution_hours": round(sum(res_hours) / len(res_hours), 1) if res_hours else None,
        "top_animal": rules.ANIMAL_LABELS[animals.most_common(1)[0][0]] if animals else None,
        "hotspots": hotspots(rows),
        "model_ready": model_ready(),
    }


@app.get("/api/export.xlsx")
def export_excel(x_admin_key: Optional[str] = Header(None)):
    """Download every report as an Excel file (dashboard "Export to Excel" button)."""
    require_admin(x_admin_key)
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    rows = [label(r) for r in query("SELECT * FROM reports ORDER BY created_at DESC")]
    wb = Workbook()
    ws = wb.active
    ws.title = "Reports"
    cols = [
        ("Report ID", "id", 11), ("Date and time", "created_at", 19), ("Animal", "animal_label", 18),
        ("Incident", "incident_label", 22), ("Urgency", "urgency", 10), ("Sent to", "agency_label", 34),
        ("Status", "status", 13), ("Latitude", "lat", 11), ("Longitude", "lng", 12),
        ("AI confidence", "confidence", 14), ("Resolved at", "resolved_at", 19), ("Sample data", "simulated", 12),
    ]
    ws.append([c[0] for c in cols])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4D3A")
        cell.alignment = Alignment(vertical="center")
    for r in rows:
        values = []
        for _, key, _w in cols:
            v = r.get(key)
            if key in ("created_at", "resolved_at") and v:
                v = datetime.fromisoformat(v)
            elif key == "status" and v:
                v = v.replace("_", " ")
            elif key == "simulated":
                v = "Yes" if v else "No"
            values.append(v)
        ws.append(values)
    for i, (_, key, width) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
        if key in ("created_at", "resolved_at"):
            for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cell.number_format = "yyyy-mm-dd hh:mm"
        if key == "confidence":
            for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cell.number_format = "0%"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = BytesIO()
    wb.save(buf)
    name = f"wildwatch-reports-{datetime.now():%Y-%m-%d-%H%M}.xlsx"
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})


@app.get("/api/health")
def health():
    """Used by Render to check the app is up."""
    return {"ok": True, "ai_model": model_ready(), "google_sheet": sheets.ENABLED}


# ---------- Pages ----------

FRONTEND = ROOT / "frontend"
app.mount("/static", StaticFiles(directory=FRONTEND), name="static")
app.mount("/uploads", StaticFiles(directory=UPLOADS), name="uploads")


@app.get("/")
def home():
    return RedirectResponse("/app")


@app.get("/manifest.webmanifest")
def manifest():
    return FileResponse(FRONTEND / "manifest.webmanifest", media_type="application/manifest+json")


@app.get("/sw.js")
def service_worker():
    # Served from the root so it can control the whole app.
    return FileResponse(FRONTEND / "sw.js", media_type="application/javascript",
                        headers={"Cache-Control": "no-cache"})


#@app.get("/report")
#def report_page():
    # The old address now opens the new app, so old home-screen icons and links still work.
 #   return RedirectResponse("/app")


#@app.get("/report-old")
#def old_report_page():
    # The previous resident page, kept as a backup.
    #return FileResponse(FRONTEND / "report.html")


@app.get("/dashboard")
def dashboard_page():
    return FileResponse(FRONTEND / "admin.html")


#@app.get("/dashboard-old")
#def old_dashboard_page():
    # The previous dashboard, kept as a backup.
 #   return FileResponse(FRONTEND / "dashboard.html")


# ---------- New resident app (Map and Tips use these two public endpoints) ----------

@app.get("/app")
def new_app_page():
    return FileResponse(FRONTEND / "app.html")


@app.get("/api/public/sightings")
def public_sightings(days: int = 14):
    """Recent sightings for the residents' map. Locations are rounded to about 100 m
    and nothing personal is included (no photo, note or exact spot)."""
    days = max(1, min(days, 60))
    since = datetime.fromtimestamp(datetime.now().timestamp() - days * 86400).isoformat(timespec="seconds")
    rows = query("SELECT lat, lng, animal, incident, created_at FROM reports WHERE created_at >= ? "
                 "ORDER BY created_at DESC LIMIT 1500", (since,))
    return [{"lat": round(r["lat"], 3), "lng": round(r["lng"], 3), "animal": r["animal"],
             "incident": r["incident"], "day": r["created_at"][:10]} for r in rows]


@app.get("/api/public/tips")
def public_tips():
    """Safety tips for every animal, so the app and the rules always say the same thing."""
    return {"labels": rules.ANIMAL_LABELS, "tips": rules.TIPS}
